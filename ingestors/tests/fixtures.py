#!/usr/bin/env python3
"""
Arquivos de exemplo dos drivers de ingestão, gerados em código (#50) — nada binário versionado.

Cada arquivo leva uma PALAVRA-ÂNCORA própria (um topônimo que não aparece em nenhum outro),
para o teste ponta a ponta provar que a busca acha o conteúdo que veio por aquele driver.

    python3 fixtures.py --out DIR     # grava os exemplos em DIR (usado por tools/e2e_ingest.sh)

pdf e pptx saem só com a biblioteca padrão; xlsx e docx precisam de openpyxl e python-docx
(as mesmas dependências dos drivers) — sem elas, aquele exemplo não é gerado.
"""
import os
import sys
# Rodando de dentro de ingestors/, o diretório entra no sys.path e `import docx`/`import csv`
# achariam os DRIVERS (docx.py, csv.py) em vez das bibliotecas — o mesmo sombreamento que cada
# driver evita (ver ingestors/README.md). Tira a pasta dos drivers do caminho antes de importar.
_drivers = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path[:] = [p for p in sys.path if os.path.abspath(p or os.getcwd()) != _drivers]
import io
import zipfile

ANCORAS = {
    "csv": "Quixeramobim",
    "xlsx": "Itapecerica",
    "docx": "Pindamonhangaba",
    "pptx": "Paranapiacaba",
    "pdf": "Mogimirim",
}


# ───────────────────────────── csv ─────────────────────────────
def csv_bytes():
    """Ponto e vírgula + BOM (o que o Excel brasileiro exporta): o driver normaliza para vírgula."""
    linhas = ["cidade;estado;habitantes", f"{ANCORAS['csv']};CE;80000", "Crateús;CE;75000"]
    return ("﻿" + "\n".join(linhas) + "\n").encode("utf-8")


# ───────────────────────────── xlsx ─────────────────────────────
def xlsx_bytes():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cidades"
    ws.append(["cidade", "estado", "habitantes"])
    ws.append([ANCORAS["xlsx"], "SP", 50000])
    ws.append([None, None, None])            # linha vazia: o driver descarta
    wb.create_sheet("Vazia")                 # aba sem conteúdo: o driver pula
    ws2 = wb.create_sheet("Notas")
    ws2.append(["observação"])
    ws2.append(["segunda aba com conteúdo"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ───────────────────────────── docx ─────────────────────────────
def docx_bytes():
    import docx
    d = docx.Document()
    d.add_heading("Relatório de viagem", level=1)
    d.add_paragraph(f"A comitiva passou por {ANCORAS['docx']} na terça-feira.")
    d.add_paragraph("")                      # parágrafo vazio: o driver descarta
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "item", "valor"
    t.cell(1, 0).text, t.cell(1, 1).text = "diária", "350"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ───────────────────────────── pptx ─────────────────────────────
_NS = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
       'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
       'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')
_RELS_NS = 'xmlns="http://schemas.openxmlformats.org/package/2006/relationships"'
_T_SLIDE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"
_T_NOTES = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesSlide"


def _sp(*paragrafos):
    """Uma caixa de texto; cada parágrafo é uma lista de runs (o PowerPoint pica a frase em
    runs quando muda a formatação no meio)."""
    ps = "".join("<a:p>" + "".join(f"<a:r><a:t>{r}</a:t></a:r>" for r in runs) + "</a:p>"
                 for runs in paragrafos)
    return f"<p:sp><p:txBody>{ps}</p:txBody></p:sp>"


def _tabela(linhas):
    trs = "".join("<a:tr>" + "".join(f"<a:tc><a:txBody><a:p><a:r><a:t>{c}</a:t></a:r></a:p></a:txBody></a:tc>"
                                     for c in linha) + "</a:tr>" for linha in linhas)
    return f"<p:graphicFrame><a:graphic><a:graphicData><a:tbl>{trs}</a:tbl></a:graphicData></a:graphic></p:graphicFrame>"


def _slide(*formas):
    return f'<?xml version="1.0" encoding="UTF-8"?><p:sld {_NS}><p:cSld><p:spTree>{"".join(formas)}</p:spTree></p:cSld></p:sld>'


def pptx_bytes(com_ordem=True):
    """Dois slides cuja ORDEM declarada (sldIdLst) é a inversa dos nomes dos arquivos: slide2.xml
    vem primeiro. Com `com_ordem=False` o presentation.xml some e vale a ordem dos nomes."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
        if com_ordem:
            z.writestr("ppt/presentation.xml",
                       f'<?xml version="1.0"?><p:presentation {_NS}><p:sldIdLst>'
                       '<p:sldId id="256" r:id="rId2"/><p:sldId id="257" r:id="rId1"/></p:sldIdLst></p:presentation>')
            z.writestr("ppt/_rels/presentation.xml.rels",
                       f'<?xml version="1.0"?><Relationships {_RELS_NS}>'
                       f'<Relationship Id="rId1" Type="{_T_SLIDE}" Target="slides/slide1.xml"/>'
                       f'<Relationship Id="rId2" Type="{_T_SLIDE}" Target="slides/slide2.xml"/></Relationships>')
        # slide1 (o SEGUNDO na apresentação): tabela + nota do apresentador
        z.writestr("ppt/slides/slide1.xml", _slide(_sp(["Custos"]), _tabela([["item", "valor"], ["trem", "42"]])))
        z.writestr("ppt/slides/_rels/slide1.xml.rels",
                   f'<?xml version="1.0"?><Relationships {_RELS_NS}>'
                   f'<Relationship Id="rId1" Type="{_T_NOTES}" Target="../notesSlides/notesSlide1.xml"/></Relationships>')
        z.writestr("ppt/notesSlides/notesSlide1.xml",
                   f'<?xml version="1.0"?><p:notes {_NS}><p:cSld><p:spTree>{_sp(["Lembrar do horário"])}</p:spTree></p:cSld></p:notes>')
        # slide2 (o PRIMEIRO): frase picada em runs — tem que sair colada
        z.writestr("ppt/slides/slide2.xml",
                   _slide(_sp(["Visita a ", ANCORAS["pptx"][:5], ANCORAS["pptx"][5:]], ["Roteiro da serra"])))
    return buf.getvalue()


# ───────────────────────────── pdf ─────────────────────────────
def _pdf_texto(s):
    """Texto num literal PDF: escapa ( ) \\ e codifica em cp1252 (WinAnsiEncoding cobre o português)."""
    return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)").encode("cp1252")


def pdf_bytes(paginas=None):
    """PDF mínimo e válido, montado à mão (xref com offsets corretos). `paginas` = lista de
    listas de linhas; página com lista vazia não tem texto nenhum (simula um scan sem OCR)."""
    if paginas is None:
        paginas = [[f"Relatório da unidade de {ANCORAS['pdf']}", "Produção de café em alta"],
                   ["Segunda página: conclusão e próximos passos"]]
    objs = []                                   # (número, bytes do corpo)
    n_pag = len(paginas)
    # 1 catálogo · 2 árvore de páginas · 3 fonte · depois, por página: página + conteúdo
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(n_pag))
    objs.append((1, b"<< /Type /Catalog /Pages 2 0 R >>"))
    objs.append((2, f"<< /Type /Pages /Kids [{kids}] /Count {n_pag} >>".encode()))
    objs.append((3, b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"))
    for i, linhas in enumerate(paginas):
        pag, cont = 4 + 2 * i, 5 + 2 * i
        corpo = b"BT /F1 12 Tf 72 720 Td 16 TL " + b" ".join(b"(" + _pdf_texto(l) + b") Tj T*" for l in linhas) + b" ET"
        objs.append((pag, f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                          f"/Resources << /Font << /F1 3 0 R >> >> /Contents {cont} 0 R >>".encode()))
        objs.append((cont, b"<< /Length " + str(len(corpo)).encode() + b" >>\nstream\n" + corpo + b"\nendstream"))
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = {}
    for num, corpo in sorted(objs):
        offsets[num] = out.tell()
        out.write(f"{num} 0 obj\n".encode() + corpo + b"\nendobj\n")
    xref = out.tell()
    total = max(offsets) + 1
    out.write(f"xref\n0 {total}\n0000000000 65535 f \n".encode())
    for num in range(1, total):
        out.write(f"{offsets[num]:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {total} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


# ───────────────────────────── gravação ─────────────────────────────
GERADORES = {"csv": csv_bytes, "xlsx": xlsx_bytes, "docx": docx_bytes, "pptx": pptx_bytes, "pdf": pdf_bytes}


def grava(pasta):
    """Grava `exemplo.<fmt>` para cada formato que der para gerar aqui; devolve os gravados."""
    os.makedirs(pasta, exist_ok=True)
    feitos = []
    for fmt, gerar in GERADORES.items():
        try:
            dados = gerar()
        except ImportError as e:
            sys.stderr.write(f"fixtures: {fmt} não gerado ({e.name} ausente)\n")
            continue
        with open(os.path.join(pasta, f"exemplo.{fmt}"), "wb") as f:
            f.write(dados)
        feitos.append(fmt)
    return feitos


if __name__ == "__main__":
    if len(sys.argv) != 3 or sys.argv[1] != "--out":
        sys.stderr.write("uso: python3 fixtures.py --out DIR\n")
        sys.exit(2)
    print(" ".join(grava(sys.argv[2])))
